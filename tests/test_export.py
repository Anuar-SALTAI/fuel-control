from io import BytesIO

from openpyxl import load_workbook

from fuel_control.export import records_to_excel


def test_excel_export():
    record = {key: value for key, value in zip(
        ["month", "vehicle_name", "season", "norm", "mileage", "filled", "opening_balance",
         "closing_balance", "normative_consumption", "actual_consumption", "difference"],
        ["2026-09", "Газель", "summer", 26, 100, 26, 0, 0, 26, 26, 0],
    )}
    content = records_to_excel([record], {"month": "Ай", "sheet": "Тарих"})
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == ["Тарих"]
    assert workbook["Тарих"]["A2"].value == "2026-09"
